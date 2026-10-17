"""TSX's own list of mining issuers, read as next quarter's source."""
import os
import sys
import tempfile
import unittest

import openpyxl

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(HERE, "fixtures"))

from mmparser.cli import to_payload  # noqa: E402
from mmparser.ingest import ingest  # noqa: E402
from build_tsx_fixture import build  # noqa: E402


class TsxList(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = build(os.path.join(self.dir.name, "tsx-list.xlsx"))

    def tearDown(self):
        self.dir.cleanup()

    def parse(self, path=None):
        companies, report = ingest(path or self.path)
        return to_payload(companies, report)

    def test_recognised_by_its_sheets_and_read_whole(self):
        p = self.parse()
        self.assertEqual(p["kind"], "tsx_list", "only issuer-list sheets, so not a whitespace workbook")
        self.assertEqual(p["period"], "2026-08-31", "a month written out still dates the period")
        self.assertEqual(p["report"]["blocking"], [])
        # Every issuer, the C$12M one too: the threshold is the application's.
        self.assertEqual(sorted(c["ticker"] for c in p["companies"]),
                         ["CARY", "DEVL", "EAST", "EDGE", "EXPL", "NEWC", "NRTH", "ROYL", "SOCO", "TINY", "WSTC"])
        totals = p["report"]["totals"]
        self.assertEqual(len(totals), 4, "a count and a market-cap total for each sheet")
        for label, (actual, expected) in totals.items():
            self.assertEqual(actual, expected, label)

    def test_each_issuer_carries_what_the_period_needs(self):
        by = {c["ticker"]: c for c in self.parse()["companies"]}
        n = by["NRTH"]
        self.assertEqual(n["exchange"], "TSX")
        self.assertEqual(n["co_id"], "C-001")
        self.assertEqual(n["head_office_region"], "Ontario, Canada")
        self.assertEqual(n["regions"]["CANADA"], ["ON", "QC"])
        self.assertEqual(n["regions"]["LATIN AMERICA"], ["Peru"])
        self.assertEqual(sorted(n["commodities"]), ["Copper", "Gold"])
        self.assertTrue(n["venture_graduate"])
        self.assertEqual(by["WSTC"]["head_office_region"], "Nevada, United States")
        self.assertEqual(by["EAST"]["head_office_region"], "Chile")
        self.assertEqual(by["EXPL"]["commodities"], ["Cobalt", "Graphite"], "Other Properties are commodities")
        self.assertFalse(by["EXPL"]["venture_graduate"], "the TSXV sheet has no graduate column")

    def test_tsx_initials_for_a_country_are_read(self):
        from mmparser.aliases import normalise_head_office
        self.assertEqual(normalise_head_office("UAE", "Other"), ("United Arab Emirates", None))

    def test_a_sheet_that_does_not_tie_to_its_own_totals_blocks(self):
        wb = openpyxl.load_workbook(self.path)
        wb["TSX MM Issuers August 2026"]["C8"] = 7   # six issuers on the sheet
        bad = os.path.join(self.dir.name, "short.xlsx")
        wb.save(bad)
        blocking = self.parse(bad)["report"]["blocking"]
        self.assertTrue(any(f["code"] == "proof_total" for f in blocking), blocking)

    def test_an_issuer_listed_twice_blocks(self):
        wb = openpyxl.load_workbook(self.path)
        ws = wb["TSXV MM Issuers August 2026"]
        ws.cell(row=11, column=3, value="TSX")      # Eastco, now claiming the TSX listing too
        ws.cell(row=11, column=5, value="NRTH")
        bad = os.path.join(self.dir.name, "twice.xlsx")
        wb.save(bad)
        blocking = self.parse(bad)["report"]["blocking"]
        self.assertTrue(any(f["code"] == "duplicate_issuer" for f in blocking), blocking)

    def test_the_whitespace_workbook_is_still_a_workbook(self):
        fixture = os.path.join(HERE, "fixtures", "synthetic_whitespace.xlsx")
        self.assertEqual(self.parse(fixture)["kind"], "workbook")


if __name__ == "__main__":
    unittest.main()
