"""The export as the web application builds it: a payload in, a workbook out.

The command line reads SQLite and calls build_workbook. The application reads
its own database -- Postgres in production, where this Python cannot go -- and
hands the same rows over as JSON. These tests hold the two paths together: the
service must build the SAME workbook from the payload that read_period would
have produced, and must not fall over on a payload missing optional keys.
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from mmparser.export import read_period  # noqa: E402
from mmparser.export_service import flat_csv_text, normalise, workbook_bytes  # noqa: E402
from mmparser.service import Refused  # noqa: E402

FIXTURE = os.path.join(ROOT, "tests", "fixtures", "synthetic_whitespace.xlsx")


def _jsonable(data):
    """The payload as it crosses the wire: sets become lists, tuples arrays."""
    return json.loads(json.dumps(data, default=lambda o: sorted(o) if isinstance(o, set) else str(o)))


@unittest.skipIf(shutil.which("node") is None, "node commits the period this reads")
class ExportService(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="mm-export-service-")
        payload = os.path.join(cls.tmp, "period.json")
        db = os.path.join(cls.tmp, "period.db")
        subprocess.check_call(
            [sys.executable, "-m", "mmparser.cli", FIXTURE, "--json", payload],
            cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        subprocess.check_call(
            ["node", "scripts/commit_period.mjs", payload, db, "--fresh"],
            cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        cls.data = _jsonable(read_period(db))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_the_payload_builds_the_same_workbook_the_command_line_writes(self):
        wb = load_workbook(io.BytesIO(workbook_bytes(self.data)))
        self.assertIn("A.02 Matrix", wb.sheetnames)
        self.assertIn("Cover", wb.sheetnames)
        self.assertIn("Provenance", wb.sheetnames)
        matrix = wb["A.02 Matrix"]
        names = [matrix.cell(row=r, column=4).value for r in range(6, 6 + len(self.data["companies"]))]
        self.assertEqual([c["name"] for c in self.data["companies"]], names)
        # The stages arrive as a list over the wire and must still mark cells.
        marked = [matrix.cell(row=r, column=c).value
                  for r in range(6, 6 + len(self.data["companies"])) for c in range(16, 20)]
        self.assertIn("X", marked, "no stage column was written")

    def test_a_payload_missing_optional_keys_is_a_blank_cell_not_a_500(self):
        bare = {"period": {"label": "Q3-2026 (2026-06-30)"},
                "companies": [{"name": "Northco Mining Corp."}]}
        wb = load_workbook(io.BytesIO(workbook_bytes(bare)))
        self.assertEqual(wb["A.02 Matrix"].cell(row=6, column=4).value, "Northco Mining Corp.")

    def test_a_payload_that_is_not_a_period_is_refused_by_name(self):
        for bad in [None, "a period", {"period": {}}, {"companies": []}]:
            with self.assertRaises(Refused):
                normalise(bad)
        with self.assertRaises(Refused) as caught:
            normalise({"period": {}, "companies": [{}] * 10_001})
        self.assertEqual(caught.exception.status, 413)

    def test_the_flat_csv_carries_the_notices_and_the_companies(self):
        text = flat_csv_text(self.data)
        self.assertTrue(text.startswith("# "), "the licence notices lead the file")
        self.assertIn(self.data["companies"][0]["name"], text)

    def test_the_module_runs_as_a_command_reading_stdin(self):
        out = subprocess.run([sys.executable, "-m", "mmparser.export_service"],
                             cwd=ROOT, input=json.dumps(self.data).encode("utf-8"),
                             stdout=subprocess.PIPE, check=True).stdout
        self.assertEqual(out[:2], b"PK")
        self.assertIn("A.02 Matrix", load_workbook(io.BytesIO(out)).sheetnames)


if __name__ == "__main__":
    unittest.main()


def _payload_with_fees():
    """Two companies, two currencies -- what Run 1 actually produced."""
    return {
        "period": {"label": "Q3-2026 (2026-05-31)", "market_cap_as_of": "2026-05-31",
                   "threshold_amount": 200000000, "threshold_operator": "gte",
                   "threshold_currency": "CAD", "proximity_band_pct": 2, "status": "published",
                   "revision": 6, "rule_set_version": "1.0.0"},
        "companies": [
            {"name": "Agnico Eagle Mines Limited", "tier": 1, "status": "classified",
             "footprint": "canada_and_abroad", "rule_set_version": "1.0.0",
             "stages": ["production", "exploration", "development"],
             "values": {"root_ticker": "AEM", "exchange": "TSX", "market_cap_cad": 78000000000,
                        "auditor": "Ernst & Young", "head_office_location": "Toronto",
                        "audit_fee": {"amount": 8052000, "currency": "CAD", "fiscal_year": 2025},
                        "tax_fee": {"amount": 382000, "currency": "CAD", "fiscal_year": 2025}}},
            {"name": "Elemental Royalty Corporation", "tier": 4, "status": "classified",
             "footprint": "abroad", "rule_set_version": "1.0.0", "stages": ["royalty_streaming"],
             "values": {"root_ticker": "ELE", "exchange": "TSXV", "market_cap_cad": 400000000,
                        "auditor": "Davidson & Company", "head_office_location": "Littleton",
                        "audit_fee": {"amount": 517116, "currency": "USD", "fiscal_year": 2025},
                        "tax_fee": {"amount": 0, "currency": "USD", "fiscal_year": 2025}}},
        ],
        "provenance": {"source_file_sha256": "abc", "rule_set_version": "1.0.0",
                       "prior_period": None},
    }


class Presentation(unittest.TestCase):
    """What a partner sees. A workbook nobody can read is not a deliverable."""

    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(io.BytesIO(workbook_bytes(_payload_with_fees())))

    def test_every_amount_says_which_currency_it_is_in(self):
        matrix = self.wb["A.02 Matrix"]
        # Market cap is Canadian dollars, and says so rather than being a bare number.
        self.assertIn("C$", matrix.cell(row=6, column=6).number_format)
        # Fees carry the currency of the filing they came from.
        self.assertEqual(matrix.cell(row=6, column=40).value, "CAD")
        self.assertEqual(matrix.cell(row=6, column=41).value, 8052000)
        self.assertIn("C$", matrix.cell(row=6, column=41).number_format)
        self.assertEqual(matrix.cell(row=7, column=40).value, "USD")
        self.assertIn("US$", matrix.cell(row=7, column=41).number_format)
        self.assertEqual(matrix.cell(row=6, column=43).value, 2025, "the fiscal year travels too")

    def test_a_foreign_fee_is_not_converted_at_a_rate_nobody_chose(self):
        matrix = self.wb["A.02 Matrix"]
        # Canadian dollars need no conversion, so the CAD column is filled.
        self.assertEqual(matrix.cell(row=6, column=42).value, 8052000)
        # US dollars do, and there is no FX table: empty, with the reason on it.
        usd_cad = matrix.cell(row=7, column=42)
        self.assertIsNone(usd_cad.value)
        self.assertIsNotNone(usd_cad.comment, "an empty cell must say why it is empty")
        self.assertIn("USD", usd_cad.comment.text)

    def test_the_auditor_columns_are_big_four_and_everyone_else(self):
        matrix = self.wb["A.02 Matrix"]
        self.assertEqual(matrix.cell(row=6, column=37).value, "Ernst & Young")
        self.assertIsNone(matrix.cell(row=6, column=38).value)
        self.assertIsNone(matrix.cell(row=7, column=37).value)
        self.assertEqual(matrix.cell(row=7, column=38).value, "Davidson & Company")

    def test_the_tier_is_filled_and_carries_its_rule(self):
        tier = self.wb["A.02 Matrix"].cell(row=6, column=2)
        self.assertEqual(tier.value, "Tier 1")
        self.assertEqual(tier.fill.fgColor.rgb[-6:], "86BC25", "a tier reads as a conclusion")
        self.assertTrue(tier.font.bold)

    def test_every_data_sheet_is_navigable_and_printable(self):
        for title in ["A.02 Matrix", "B.01 Consol TSX - TSXV", "B.04 Auditor & Fees"]:
            ws = self.wb[title]
            self.assertTrue(ws.freeze_panes, "%s does not hold its headers" % title)
            self.assertTrue(ws.auto_filter.ref, "%s cannot be filtered" % title)
            self.assertTrue(ws.print_title_rows, "%s loses its header on page 2" % title)
            self.assertEqual(ws.page_setup.orientation, "landscape")
            self.assertEqual(ws.sheet_properties.tabColor.rgb[-6:].upper() in
                             {"86BC25", "A7CF5E", "C6E29A"}, True)

    def test_no_tab_is_empty_and_contents_says_what_each_holds(self):
        for ws in self.wb.worksheets:
            filled = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, ""))
            self.assertGreater(filled, 0, "%s has nothing on it" % ws.title)
        contents = "\n".join(str(c.value) for row in self.wb["Contents"].iter_rows() for c in row)
        for title in self.wb.sheetnames:
            if title != "Contents":
                self.assertIn(title, contents, "%s is not listed in Contents" % title)

    def test_the_provenance_says_what_the_fees_are_and_what_was_not_converted(self):
        text = "\n".join(str(c.value) for row in self.wb["Provenance"].iter_rows() for c in row)
        self.assertIn("CAD, USD", text)
        self.assertIn("unconverted", text)


class SummaryPage(unittest.TestCase):
    """The page a partner opens. Its numbers have to be the matrix's numbers."""

    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(io.BytesIO(workbook_bytes(_payload_with_fees())))
        cls.ws = cls.wb["Summary"]
        cls.cells = {c.coordinate: c.value for row in cls.ws.iter_rows() for c in row
                     if c.value is not None}

    def _formulas(self):
        return [v for v in self.cells.values() if isinstance(v, str) and v.startswith("=")]

    def test_every_range_covers_the_data_rows_and_no_others(self):
        # The defect this file exists to correct was a count that began one row
        # below the data. Two companies here: rows 6 and 7, nothing else.
        import re
        ranges = set()
        for formula in self._formulas():
            ranges.update(re.findall(r"\$[A-Z]{1,2}\$(\d+):\$[A-Z]{1,2}\$(\d+)", formula))
        self.assertTrue(ranges, "the summary has no ranges at all")
        for first, last in ranges:
            self.assertEqual((int(first), int(last)), (6, 7),
                             "a range does not cover exactly the two data rows")

    def test_every_share_is_a_share_of_the_population(self):
        shares = [v for k, v in self.cells.items()
                  if k.startswith("D") and isinstance(v, str) and v.startswith("=IF($C$")]
        self.assertGreaterEqual(len(shares), 10)
        for formula in shares:
            self.assertIn("$C$6", formula, "a share divides by something other than the total")

    def test_the_ranking_is_the_companies_deloitte_does_not_audit(self):
        names, caps = [], []
        for row in self.ws.iter_rows(min_col=2, max_col=5):
            if row[0].row > 28 and isinstance(row[0].value, str) and isinstance(row[3].value, (int, float)):
                names.append(row[0].value)
                caps.append(row[3].value)
        self.assertEqual(names, ["Agnico Eagle Mines Limited", "Elemental Royalty Corporation"])
        self.assertEqual(caps, sorted(caps, reverse=True), "largest first, or it is not a ranking")

    def test_it_says_what_research_is_still_missing(self):
        text = "\n".join(str(v) for v in self.cells.values())
        self.assertIn("no stage researched", text)
        self.assertIn("no audit fee", text)
        # And it says where "not audited by Deloitte" comes from, since the
        # firm's own client list is out of scope for this build.
        self.assertIn("not stored in this build", text)


TABLE = {
    "period": {"label": "Q3-2026 (2026-05-31)", "basis": "Published revision 9 of Q3-2026 (2026-05-31)", "revision": 9},
    "filter": "Tier: Tier 4 · Company HQ: Ontario (ON)", "count": 2, "total": 259,
    "columns": [
        {"key": "name", "label": "Company", "kind": "text"},
        {"key": "exploration", "label": "Exploration", "kind": "stage"},
        {"key": "production", "label": "Production", "kind": "stage"},
        {"key": "hq", "label": "Company HQ", "kind": "text", "field": "head_office_region"},
        {"key": "audit_fee", "label": "Audit fees", "kind": "fee", "field": "audit_fee"},
        {"key": "market_cap", "label": "Market cap (CAD)", "kind": "money"},
        {"key": "meta:commodities", "label": "Commodities", "kind": "text"},
    ],
    "rows": [
        {"name": "Northco", "exploration": "yes", "production": "unknown", "hq": "Toronto, ON",
         "audit_fee": {"amount": 1200000, "currency": "CAD", "fiscal_year": 2025},
         "market_cap": 5e9, "meta:commodities": "Gold",
         "links": {"audit_fee": "https://northco.example/circular.pdf"}},
        {"name": "=HYPERLINK(\"http://evil\")", "exploration": "not_researched", "production": "not_researched",
         "hq": None, "audit_fee": None, "market_cap": None, "links": {}},
    ],
    "sources": [{"company": "Northco", "field": "Audit fees", "value": "CAD 1200000 FY2025",
                 "kind": "Researched from a filing", "title": "Management Information Circular 2026",
                 "url": "https://northco.example/circular.pdf", "retrieved": "2026-09-12", "quote": "Audit fees: C$1,200,000"}],
}


class TableExport(unittest.TestCase):
    """The companies table as filtered on screen: a sibling of the period export."""

    def setUp(self):
        from mmparser.export_service import table_bytes
        self.wb = load_workbook(io.BytesIO(table_bytes(json.loads(json.dumps(TABLE)))))

    def test_it_is_the_table_and_its_sources(self):
        self.assertEqual(self.wb.sheetnames, ["Companies", "Sources"])

    def test_the_header_says_what_was_exported_and_how_it_was_produced(self):
        from mmparser.export import AI_NOTICE, INTERNAL_USE
        text = " ".join(str(c.value) for row in self.wb["Companies"].iter_rows(max_row=14) for c in row if c.value)
        for words in ("Published revision 9", "Tier: Tier 4", "2 of 259", INTERNAL_USE, AI_NOTICE):
            self.assertIn(words, text)

    def test_a_fee_is_three_columns_and_its_amount_links_to_its_source(self):
        ws = self.wb["Companies"]
        header = next(r for r in range(1, 30) if ws.cell(r, 2).value == "Company")
        labels = [ws.cell(header, c).value for c in range(2, 12)]
        self.assertIn("Audit fees — amount", labels)
        amount = ws.cell(header + 1, 2 + labels.index("Audit fees — amount"))
        self.assertEqual(amount.value, 1200000)
        self.assertEqual(amount.hyperlink.target, "https://northco.example/circular.pdf")

    def test_stage_marks_are_x_and_question_and_the_unresearched_are_filled_not_blank_noes(self):
        ws = self.wb["Companies"]
        header = next(r for r in range(1, 30) if ws.cell(r, 2).value == "Company")
        self.assertEqual(ws.cell(header + 1, 3).value, "X")
        self.assertEqual(ws.cell(header + 1, 4).value, "?")
        self.assertIsNone(ws.cell(header + 2, 3).value)
        self.assertEqual(ws.cell(header + 2, 3).fill.fgColor.rgb[-6:], "FFF2A8")

    def test_a_formula_typed_as_a_name_stays_text(self):
        ws = self.wb["Companies"]
        header = next(r for r in range(1, 30) if ws.cell(r, 2).value == "Company")
        cell = ws.cell(header + 2, 2)
        self.assertEqual(cell.data_type, "s", "the injection guard holds in the table export too")

    def test_the_sources_sheet_quotes_the_passage(self):
        ws = self.wb["Sources"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("Audit fees: C$1,200,000", text)
        self.assertIn("Management Information Circular 2026", text)

    def test_the_csv_leads_with_the_notices_and_the_filter_then_links_each_cited_value(self):
        from mmparser.export_service import table_csv_text
        text = table_csv_text(json.loads(json.dumps(TABLE)))
        head = [l for l in text.splitlines() if l.startswith("# ")]
        self.assertTrue(any("artificial intelligence" in l for l in head))
        self.assertIn("# Filter: Tier: Tier 4 · Company HQ: Ontario (ON)", head)
        self.assertIn("Audit fees — source", text)
        self.assertIn("https://northco.example/circular.pdf", text)
        self.assertIn("not researched", text, "a CSV has no fill, so the words say it")

    def test_a_payload_that_is_not_a_table_is_refused_by_name(self):
        from mmparser.export_service import normalise_table
        with self.assertRaises(Refused) as caught:
            normalise_table({"rows": []})
        self.assertIn("columns", caught.exception.message)
        with self.assertRaises(Refused):
            normalise_table({"columns": [{"key": "x", "label": "X", "kind": "formula"}], "rows": []})

    def test_the_function_dispatches_on_the_kind_header(self):
        with open(os.path.join(ROOT, "api", "workbook.py"), encoding="utf-8") as fh:
            src = fh.read()
        self.assertIn('self.headers.get("X-MM-Export-Kind") == "table"', src)
        self.assertIn("table_bytes(data) if table else workbook_bytes(data)", src)
