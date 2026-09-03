#!/usr/bin/env python3
"""Regression test for docs/design/00-corrected-ground-truth.md.

Every number asserted in the design document is re-measured here against the
real workbooks. Run before publishing any revision of the design.

Usage:  python3 scripts/verify_ground_truth.py [--reference-dir DIR]

The workbooks are NOT in version control (licence + personal data, see
docs/design/11-security-privacy-compliance.md), so this runs locally or in a
nightly job with access to the private copies -- never in CI.
"""
import argparse, os, re, sys, zipfile, warnings
warnings.filterwarnings("ignore")
import openpyxl

WB_NAME = "M&M - Whitespace Analysis Q3-2026.xlsx"
LSEG_NAME = "Canada Mining Company Screen 4-30-2026.xlsx"
REGIONS = ["Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF"]

results = []


def check(claim, actual, expected):
    results.append((actual == expected, claim, actual, expected))


def relative_luminance(hexcolor):
    r, g, b = (int(hexcolor[i:i + 2], 16) / 255 for i in (1, 3, 5))
    def lin(c):
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = lin(r), lin(g), lin(b)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = relative_luminance(a), relative_luminance(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def hdr(ws, row):
    return {str(ws.cell(row=row, column=c).value).replace("\n", " ").strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=row, column=c).value is not None}


def main(refdir):
    wbp = os.path.join(refdir, WB_NAME)
    lsegp = os.path.join(refdir, LSEG_NAME)
    for p in (wbp, lsegp):
        if not os.path.exists(p):
            print("MISSING: %s" % p, file=sys.stderr)
            return 2

    wbv = openpyxl.load_workbook(wbp, data_only=True)
    wbf = openpyxl.load_workbook(wbp, data_only=False)

    # --- sheets ---------------------------------------------------------
    check("workbook sheet count", len(wbv.sheetnames), 13)
    dividers = [s.title for s in wbv.worksheets
                if s.max_row <= 1 and s.sheet_state == "visible"]
    check("visible empty divider sheets", len(dividers), 2)

    m = wbv["A.02 Matrix"]
    mf = wbf["A.02 Matrix"]

    # --- population -----------------------------------------------------
    names = [m["D%d" % r].value for r in range(6, 265)]
    check("companies in matrix", sum(1 for v in names if v), 259)
    ex = [m["E%d" % r].value for r in range(6, 265)]
    check("TSX companies", sum(1 for v in ex if v == "TSX"), 144)
    check("TSXV companies", sum(1 for v in ex if v == "TSXV"), 115)
    check("other-exchange companies",
          sum(1 for v in ex if v not in ("TSX", "TSXV")), 0)

    # --- stage flags ----------------------------------------------------
    stages = {c: sum(1 for r in range(6, 265)
                     if str(m["%s%d" % (c, r)].value or "").strip())
              for c in ("P", "Q", "R", "S")}
    check("exploration flags", stages["P"], 0)
    check("development flags", stages["Q"], 0)
    check("production flags", stages["R"], 0)
    check("royalty flags", stages["S"], 17)
    blank = sum(1 for r in range(6, 265)
                if not any(str(m["%s%d" % (c, r)].value or "").strip()
                           for c in ("P", "Q", "R", "S")))
    check("companies with no stage flag", blank, 242)

    # --- footprint ------------------------------------------------------
    no_region = [r for r in range(6, 265)
                 if all(not str(m["%s%d" % (c, r)].value or "").strip()
                        for c in REGIONS)]
    check("companies with no region values", len(no_region), 12)
    check("all no-region companies are royalty-flagged",
          all(str(m["S%d" % r].value or "").strip() == "X" for r in no_region), True)
    check("no-region companies filed as 'Canada only'",
          all(m["U%d" % r].value == "Canada only" for r in no_region), True)

    # --- tiers ----------------------------------------------------------
    check("distinct tier values today",
          {m["B%d" % r].value for r in range(6, 265)}, {"4"})
    check("tier stored as text", isinstance(m["B6"].value, str), True)

    # --- market / HQ ----------------------------------------------------
    n_blank = sum(1 for r in range(6, 265) if not m["N%d" % r].value)
    n_country = sum(1 for r in range(6, 265) if m["N%d" % r].value == "Chile")
    check("blank Deloitte market", n_blank, 50)
    check("country typed into market column", n_country, 3)
    check("foreign-HQ bucket (blank + country)", n_blank + n_country, 53)

    # --- auditor --------------------------------------------------------
    ak = [m["AK%d" % r].value for r in range(6, 265)]
    check("distinct auditor values", len({v for v in ak if v is not None}), 11)
    for spelling in ("EY (Ernst & Young)", "MNP", "Deloitte ", "Other "):
        check("auditor value present: %r" % spelling, spelling in ak, True)

    # --- website --------------------------------------------------------
    ay = [m["AY%d" % r].value for r in range(6, 265)]
    populated = [v for v in ay if v not in (None, "", 0)]
    titles = [v for v in populated
              if not re.search(r"(https?://|www\.|\.(com|ca|net|org)\b)",
                               str(v), re.I)]
    check("populated website values", len(populated), 143)
    check("website values that are titles not URLs", len(titles), 45)
    check("numeric-zero sentinel present in website column", 0 in ay, True)

    # --- duplicate header names -----------------------------------------
    row5 = [mf.cell(row=5, column=c).value for c in range(1, 57)]
    dupes = {v for v in row5 if v and row5.count(v) > 1}
    check("repeated header names on the matrix", len(dupes), 4)

    # --- market cap -----------------------------------------------------
    caps = sorted(v for v in (m["F%d" % r].value for r in range(6, 265))
                  if isinstance(v, (int, float)))
    check("smallest market cap (dollars)", round(caps[0]), 201701608)
    check("all companies at or above 200M",
          sum(1 for c in caps if c >= 2e8), 259)

    # --- extracts -------------------------------------------------------
    b02 = wbv["B.02 TSX MM Issuers June 2025"]
    b03 = wbv["B.03 TSXV MM Issuers June 2025"]
    h2, h3 = hdr(b02, 10), hdr(b03, 10)
    check("Co_ID empty in TSX extract",
          sum(1 for r in range(11, b02.max_row + 1)
              if b02.cell(row=r, column=h2["Co_ID"]).value), 0)
    check("Co_ID empty in TSXV extract",
          sum(1 for r in range(11, b03.max_row + 1)
              if b03.cell(row=r, column=h3["Co_ID"]).value), 0)
    check("PO ID empty in TSXV extract",
          sum(1 for r in range(11, b03.max_row + 1)
              if b03.cell(row=r, column=h3["PO ID"]).value), 0)
    gradcol = [k for k in h2 if "Grad" in k][0]
    check("venture-graduate flags in TSX extract",
          sum(1 for r in range(11, b02.max_row + 1)
              if b02.cell(row=r, column=h2[gradcol]).value), 98)
    check("TSX extract autofilter present", b02.auto_filter.ref is not None, True)

    # --- auditor tab entity identifier ----------------------------------
    b04 = wbv["B.04 Auditor & Fees"]
    h4 = hdr(b04, 5)
    eid = [b04.cell(row=r, column=h4["Entity ID (S&P)"]).value
           for r in range(6, 265)]
    check("real entity identifiers",
          sum(1 for v in eid if isinstance(v, (int, float))), 143)
    check("entity identifier sentinels",
          sum(1 for v in eid if v == "Not found"), 116)

    # --- screener -------------------------------------------------------
    lv = openpyxl.load_workbook(lsegp, data_only=True)
    ms = lv["Mining Screen"]
    ids = [ms.cell(row=r, column=2).value for r in range(5, 1501)]
    check("screener data rows", sum(1 for v in ids if v), 1496)
    check("screener cached count is stale", ms["J5"].value, 1497)

    # --- package parts: licences and personal data ----------------------
    for path, creator, ncomments in ((wbp, "TSX Group Inc.", 3),
                                     (lsegp, "Refinitiv", 0)):
        tag = os.path.basename(path)[:10]
        z = zipfile.ZipFile(path)
        core = z.read("docProps/core.xml").decode("utf8", "replace")
        check("%s... recorded creator" % tag,
              re.search(r"<dc:creator[^>]*>(.*?)</dc:creator>", core).group(1),
              creator)
        check("%s... records a last modifier" % tag,
              bool(re.search(r"<cp:lastModifiedBy[^>]*>(.+?)</cp:lastModifiedBy>",
                             core)), True)
        cparts = [n for n in z.namelist() if re.match(r"xl/comments\d+\.xml", n)]
        total = sum(z.read(c).decode("utf8", "replace").count("<comment ")
                    for c in cparts)
        check("%s... embedded cell comments" % tag, total, ncomments)
        check("%s... contains no macros" % tag,
              not any("vbaProject" in n for n in z.namelist()), True)

    z = zipfile.ZipFile(wbp)
    comp = sum(i.compress_size for i in z.infolist())
    unc = sum(i.file_size for i in z.infolist())
    check("workbook zip expansion ratio (rounded)", round(unc / comp, 1), 8.7)

    # --- formula-injection characters already present -------------------
    at_cells = sum(1 for row in mf.iter_rows(min_row=1, max_row=mf.max_row)
                   for c in row
                   if isinstance(c.value, str) and c.value.startswith("@"))
    check("cells already beginning with '@'", at_cells, 9)

    # --- contrast -------------------------------------------------------
    check("Deloitte green on white (2dp)",
          round(contrast("#86BC25", "#FFFFFF"), 2), 2.27)
    check("Deloitte green on black (2dp)",
          round(contrast("#86BC25", "#000000"), 2), 9.23)
    check("#6B961E on white clears 3:1",
          contrast("#6B961E", "#FFFFFF") >= 3.0, True)
    check("#5E841A focus ring on white clears 3:1",
          contrast("#5E841A", "#FFFFFF") >= 3.0, True)
    check("#567C18 green text on white clears 4.5:1",
          contrast("#567C18", "#FFFFFF") >= 4.5, True)

    # --- report ---------------------------------------------------------
    width = max(len(c) for _, c, _, _ in results)
    failed = 0
    for ok, claim, actual, expected in results:
        if not ok:
            failed += 1
            print("FAIL  %-*s  actual=%r  expected=%r"
                  % (width, claim, actual, expected))
    print("\n%d/%d ground-truth assertions passed."
          % (len(results) - failed, len(results)))
    return 1 if failed else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--reference-dir", default="reference")
    sys.exit(main(ap.parse_args().reference_dir))
